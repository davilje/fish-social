from pathlib import Path
import bisect
import hashlib
import json
import math
import re
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
spec_dir = ROOT / 'docs/策划/规格/玩法'
scale_file = spec_dir / '上钩鱼尺寸与体长映射.md'
break_file = spec_dir / '钓鱼断线与断竿表现.md'
prompt_file = ROOT / 'docs/策划/提示词/玩法/fish-break-size-design-dev.prompt.md'
scale_text = scale_file.read_text()
points = [(float(a), float(b)) for a, b in re.findall(
    r'^\| (\d+\.\d+) \| (\d+) \| [\d.]+% \|', scale_text, re.M)]
assert len(points) == 18, points
assert points[0] == (0.02, 8) and points[-1] == (50, 204)
assert all(a[0] < b[0] and a[1] < b[1] for a, b in zip(points, points[1:]))

def mapped(value):
    value = min(50, max(.02, value))
    i = min(len(points) - 2, max(0, bisect.bisect_right([p[0] for p in points], value) - 1))
    a, b = points[i:i+2]
    return a[1] + (b[1] - a[1]) * math.log(value / a[0]) / math.log(b[0] / a[0])

for value, expected in [(.3, 24.35), (1.2, 50.50), (4, 93.26), (40, 192)]:
    assert abs(mapped(value) - expected) < .005
samples = [mapped(.02 * (50/.02)**(i/10000)) for i in range(10001)]
assert all(a <= b for a, b in zip(samples, samples[1:]))
for length, expected in points[1:-1]:
    assert abs(mapped(length - 1e-9) - mapped(length + 1e-9)) < .00001

links = 0
for path in [scale_file, break_file, prompt_file,
             spec_dir / 'README.md', ROOT / 'docs/策划/提示词/玩法/README.md']:
    for target in re.findall(r'\]\(([^)]+)\)', path.read_text()):
        if '://' in target or target.startswith('#'):
            continue
        assert (path.parent / target.split('#')[0]).exists(), (path, target)
        links += 1

before = openpyxl.load_workbook(OUT / 'plan-before.xlsx')
after = openpyxl.load_workbook(ROOT / '项目开发需求计划表.xlsx')
assert before.sheetnames == after.sheetnames
ids = {'FISH-BREAK-01', 'FISH-SCALE-01'}
plan_rows = []
for sa, sb in zip(before, after):
    old = list(sa.iter_rows(values_only=True))
    new = list(sb.iter_rows(values_only=True))
    if sb == after.worksheets[0]:
        additions = [(i, row) for i, row in enumerate(new, 1) if len(row) > 1 and row[1] in ids]
        assert {r[1] for i, r in additions} == ids and len(additions) == 2
        assert [r for r in new if len(r) < 2 or r[1] not in ids] == old
        for row_number, row in additions:
            assert row[0] == '已确认' and row[9] == '2026-10-31' and not row[10]
            assert (ROOT / row[8]).exists()
            assert sb.cell(row_number, 8).alignment.wrap_text
            assert sb.row_dimensions[row_number].height >= 70
            plan_rows.append({'id': row[1], 'sheet': sb.title, 'row': row_number})
    else:
        assert old == new, sa.title

for filename in ['项目开发需求计划表.xlsx', '策划进度看板.html']:
    assert (ROOT / filename).read_bytes() == (ROOT / 'docs/策划' / filename).read_bytes()
board = (ROOT / '策划进度看板.html').read_text()
assert all(i in board for i in ids)
report = {'result': 'passed', 'mapping_control_points': points,
          'interpolation_and_continuity_samples': 10001,
          'local_markdown_links_checked': links, 'plan_rows': plan_rows,
          'existing_workbook_values_preserved': True,
          'canonical_workbook_and_board_copies_equal': True,
          'game_implementation': 'not changed; design only',
          'visual_tuning': 'proposed; game preview and Windows acceptance pending'}
(OUT / 'verification.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n')
print(json.dumps(report, ensure_ascii=False))
